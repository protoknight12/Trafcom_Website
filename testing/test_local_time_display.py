"""
pytest: timestamps are stored as naive UTC and shown in the app's zone (Europe/Sofia).
Guards to_local() (+2 h in winter, +3 h in summer, None/date untouched) and that the
activity-log page and its xlsx export print the local time, not the stored UTC one.

Run with:
    pytest testing/test_local_time_display.py -v
"""
import io
import os
import tempfile
from datetime import date, datetime, timezone

_db_fd, _db_path = tempfile.mkstemp(suffix='.db')
os.close(_db_fd)
os.environ['SECRET_KEY'] = 'test-secret-key-not-for-production'
os.environ['DATABASE_URL'] = f'sqlite:///{_db_path}'

import pytest
from openpyxl import load_workbook
from werkzeug.security import generate_password_hash

from app import app as flask_app, db, User, ActivityLog, to_local, limiter


def test_to_local_winter_and_summer():
    assert to_local(datetime(2026, 1, 15, 10, 0)).replace(tzinfo=None) == datetime(2026, 1, 15, 12, 0)
    assert to_local(datetime(2026, 7, 15, 10, 0)).replace(tzinfo=None) == datetime(2026, 7, 15, 13, 0)


def test_to_local_midnight_flips_the_date():
    assert to_local(datetime(2026, 7, 15, 22, 30)).date() == date(2026, 7, 16)


def test_to_local_aware_and_passthrough():
    assert to_local(datetime(2026, 7, 15, 10, 0, tzinfo=timezone.utc)).hour == 13
    assert to_local(None) is None
    d = date(2026, 7, 15)
    assert to_local(d) is d


@pytest.fixture
def admin_client():
    flask_app.config.update(TESTING=True, WTF_CSRF_ENABLED=False)
    limiter.reset()
    with flask_app.app_context():
        db.create_all()
        db.session.add(User(username='tz_admin', password=generate_password_hash('irrelevant123'), role='admin'))
        db.session.add(ActivityLog(timestamp=datetime(2026, 7, 15, 10, 0, 0), username='someone', role='admin', action='probe-action'))
        db.session.commit()
        c = flask_app.test_client()
        c.post('/login', data={'username': 'tz_admin', 'password': 'irrelevant123'})
        yield c
        db.session.remove()
        db.drop_all()


def test_activity_log_page_shows_local_time(admin_client):
    html = admin_client.get('/admin/log').get_data(as_text=True)
    assert '2026-07-15 13:00:00' in html
    assert '2026-07-15 10:00:00' not in html


def test_activity_log_export_shows_local_time(admin_client):
    r = admin_client.get('/admin/log/export')
    assert r.status_code == 200
    ws = load_workbook(io.BytesIO(r.data)).active
    cells = [c.value for row in ws.iter_rows() for c in row]
    assert '2026-07-15 13:00:00' in cells
